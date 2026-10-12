import unittest
import csv
import io
from openpyxl import Workbook
import test_students
from models.student import Student, StudentLogin
from models.user import UserRole
from routers.super_admin import router
from utils.passwords import verify_password


class StudentBulkTests(unittest.TestCase):
    tearDown = test_students.StudentTests.tearDown
    student_payload = test_students.StudentTests.student_payload
    def setUp(self):
        test_students.StudentTests.setUp(self)
        self.client.app.include_router(router)
        data = self.student_payload()
        self.params = dict(school_id=1, **{key: data[key] for key in (
            'class_id', 'section_id', 'session_id', 'caste_category_id', 'fee_category_id')})

    def upload(self, text):
        workbook = Workbook()
        for row in csv.reader(io.StringIO(text)):
            workbook.active.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return self.client.post('/super-admin/students/bulk-upload', params=self.params,
            files={'file': ('students.xlsx', stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})

    def test_bulk_import_and_shared_login(self):
        response = self.upload('admission_number,name,father_name,mobile_number\n001,Alice,Father,0123456789\n002,Bob,Father,0123456789\n')
        self.assertEqual(response.status_code, 201, response.text)
        self.assertEqual(response.json()['count'], 2)
        students = self.db.query(Student).order_by(Student.id).all()
        self.assertEqual([s.admission_number for s in students], ['001', '002'])
        self.assertEqual(students[0].login_id, students[1].login_id)
        self.assertTrue(verify_password('0123456789', self.db.query(StudentLogin).one().password))
        self.assertEqual(students[0].session_id, self.params['session_id'])
        self.assertEqual(self.upload('admission_number,name,father_name,mobile_number\n001,C,F,9876543210').json()['failed'], 1)
        self.assertEqual(self.db.query(Student).count(), 2)

    def test_invalid_rows_reported_and_valid_rows_inserted(self):
        response = self.upload('admission_number,name,father_name,mobile_number,house_id\n001,Valid,Father,123,\n002,,Father,456,\n003,Wrong house,Father,789,99999')
        self.assertEqual(response.status_code, 201, response.text)
        body = response.json()
        self.assertEqual((body['total'], body['inserted'], body['failed']), (3, 1, 2))
        self.assertEqual([row['row'] for row in body['failed_students']], [3, 4])
        self.assertEqual(body['failed_students'][0]['mobile_number'], '456')
        self.assertIn('errors', body['failed_students'][0])
        self.assertIn('house_id', body['failed_students'][1]['reason'])
        self.assertEqual(self.db.query(Student).count(), 1)

    def test_categories_in_csv_and_transaction_rollback(self):
        from unittest.mock import patch
        from sqlalchemy.exc import IntegrityError
        caste = self.params.pop('caste_category_id')
        fee = self.params.pop('fee_category_id')
        text = ('admission_number,name,father_name,mobile_number,caste_category_id,fee_category_id\n'
                f'001,A,F,123,{caste},{fee}\n')
        with patch.object(self.db, 'commit', side_effect=IntegrityError('insert', {}, Exception())):
            self.assertEqual(self.upload(text).status_code, 409)
        self.assertEqual(self.db.query(Student).count(), 0)
        self.assertEqual(self.db.query(StudentLogin).count(), 0)
        self.assertEqual(self.upload(text).status_code, 201)

    def test_excel_numeric_cells_and_formula_rejection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['admission_number', 'name', 'father_name', 'mobile_number'])
        sheet.append([101, 'Alice', 'Father', 9876543210])
        def send():
            stream = io.BytesIO()
            workbook.save(stream)
            return self.client.post('/super-admin/students/bulk-upload', params=self.params,
                files={'file': ('students.xlsx', stream.getvalue())})
        self.assertEqual(send().status_code, 201)
        self.assertEqual(self.db.query(Student).one().mobile_number, '9876543210')
        sheet['A2'] = 102
        sheet['B2'] = '=1+1'
        self.assertEqual(send().json()['failed'], 1)
        self.assertEqual(self.db.query(Student).count(), 1)
        workbook.close()

    def test_rejects_invalid_excel(self):
        for filename in ('students.csv', 'students.xlsx'):
            response = self.client.post('/super-admin/students/bulk-upload', params=self.params,
                files={'file': (filename, b'not an Excel workbook')})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.db.query(Student).count(), 0)

    def test_categories_are_optional(self):
        self.params.pop('caste_category_id')
        self.params.pop('fee_category_id')
        response = self.upload('admission_number,name,father_name,mobile_number\n001,Alice,Father,0123456789')
        self.assertEqual(response.status_code, 201, response.text)
        student = self.db.query(Student).one()
        self.assertIsNone(student.caste_category_id)
        self.assertIsNone(student.fee_category_id)
        response = self.client.get(f'/schools/1/students/{student.id}')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['data']['student_info']['caste_category_id'], '')

    def test_full_student_fields_are_saved(self):
        response = self.upload('admission_number,first_name,father_name,mobile_number,last_name,admission_date,father_date_of_birth,weight_kg,present_address.line_1,present_address.city,permanent_address.line_1,previous_school.school_name\nADM99,Alice,Father,0123456789,Kumar,2024-06-01,1980-01-01,25.5,Present Street,Delhi,Permanent Street,Old School')
        self.assertEqual(response.status_code, 201, response.text)
        item = self.db.query(Student).one()
        self.assertEqual(item.last_name, 'Kumar')
        self.assertEqual(str(item.admission_date), '2024-06-01')
        self.assertEqual(str(item.father_dob), '1980-01-01')
        self.assertEqual(item.weight_kg, 25.5)
        addresses = {a.address_type: a for a in item.addresses}
        self.assertEqual(addresses['present'].line_1, 'Present Street')
        self.assertEqual(addresses['permanent'].line_1, 'Permanent Street')
        self.assertEqual(item.previous_school.school_name, 'Old School')

    def test_excel_scope_must_match_query(self):
        response = self.upload('admission_number,first_name,father_name,mobile_number,session_id\n001,A,F,123,99999')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['failed'], 1)
        self.assertEqual(self.db.query(Student).count(), 0)

    def test_excel_admission_survives_school_api_update(self):
        from schemas.student_bulk import StudentBulkRow, bulk_payload
        response = self.upload('admission_number,first_name,father_name,mobile_number\n0007-EXCEL,Alice,Father,0123456789')
        self.assertEqual(response.status_code, 201, response.text)
        student_id = response.json()['data'][0]['id']
        item = self.db.query(Student).one()
        self.assertEqual(item.admission_number, '0007-EXCEL')
        self.assertIsNone(item.admission_sequence)
        payload = bulk_payload(StudentBulkRow(admission_number='0007-EXCEL', first_name='Alice',
            father_name='Father', mobile_number='0123456789'),
            self.params['session_id'], self.params['class_id'], self.params['section_id'],
            self.params['caste_category_id'], self.params['fee_category_id'])
        response = self.client.put('/schools/students',
            params={'school_id': 1, 'student_id': student_id}, json=payload.model_dump(mode='json'))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['data']['student_info']['admission_number'], '0007-EXCEL')
        self.assertEqual(self.db.query(StudentLogin).count(), 1)

    def test_duplicates_skipped_and_other_rows_inserted(self):
        header = 'admission_number,first_name,father_name,mobile_number\n'
        self.assertEqual(self.upload(header + '001,Original,Father,123').status_code, 201)
        response = self.upload(header + '001,,,\n002,New,Father,456\n002,Duplicate,Father,789\n003,Third,Father,456')
        self.assertEqual(response.status_code, 201, response.text)
        body = response.json()
        self.assertEqual((body['total'], body['inserted'], body['failed']), (4, 2, 2))
        self.assertEqual([r['row'] for r in body['failed_rows']], [2, 4])
        students = self.db.query(Student).order_by(Student.id).all()
        self.assertEqual([s.first_name for s in students], ['Original', 'New', 'Third'])
        self.assertEqual(students[1].login_id, students[2].login_id)
        response = self.upload(header + '001,,,')
        self.assertEqual(response.status_code, 201)
        self.assertEqual((response.json()['inserted'], response.json()['failed']), (0, 1))

    def test_list_includes_incomplete_imported_profiles(self):
        self.params.pop('caste_category_id')
        self.params.pop('fee_category_id')
        self.upload('admission_number,first_name,father_name,mobile_number\n001,Alice,Father,123\n002,Bob,Father,456')
        item = self.db.query(Student).first()
        for field in ('gender', 'nationality', 'mother_name', 'line_1', 'city', 'state', 'country'):
            setattr(item, field, '')
        item.addresses.clear()
        self.db.commit()
        response = self.client.get('/schools/students', params={
            'school_id': 1, 'session_id': self.params['session_id'],
            'status': 'true', 'offset': 0, 'limit': 50})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(response.json()['data']), 2)
        data = response.json()['data'][0]
        self.assertEqual(data['student_info']['date_of_birth'], '')
        self.assertEqual(data['student_info']['caste_category_id'], '')
        self.assertEqual(data['present_address']['line_1'], '')
        self.assertEqual(data['parent_info']['mother_name'], '')
        self.assertEqual(data['previous_school'], '')
        self.assertIsInstance(data['id'], int)

    def test_read_schema_accepts_null_student_details(self):
        from schemas.student import StudentInfo, ParentInfo, StudentAddress, StudentInfoResponse, ParentInfoResponse, StudentAddressResponse
        from pydantic import ValidationError
        for write_model, read_model in ((StudentInfo, StudentInfoResponse),
                                        (ParentInfo, ParentInfoResponse),
                                        (StudentAddress, StudentAddressResponse)):
            values = dict.fromkeys(write_model.model_fields, None)
            parsed = read_model.model_validate(values)
            self.assertTrue(all(value is None for value in parsed.model_dump().values()))
            with self.assertRaises(ValidationError):
                write_model.model_validate(values)

    def test_database_row_failure_does_not_block_other_students(self):
        from unittest.mock import patch
        from sqlalchemy.exc import IntegrityError
        from routers.students import update_student_details
        def save(db, school_id, item, payload):
            update_student_details(db, school_id, item, payload)
            if item.admission_number == '002':
                raise IntegrityError('insert', {}, Exception())
        with patch('routers.student_bulk.update_student_details', side_effect=save):
            response = self.upload('admission_number,name,father_name,mobile_number\n001,A,F,123\n002,B,F,456\n003,C,F,789')
        self.assertEqual(response.status_code, 201, response.text)
        self.assertEqual((response.json()['inserted'], response.json()['failed']), (2, 1))
        self.assertEqual(response.json()['failed_students'][0]['name'], 'B')
        self.assertEqual([s.admission_number for s in self.db.query(Student).order_by(Student.id)], ['001', '003'])
        self.assertEqual(self.db.query(StudentLogin).count(), 2)

    def test_download_failed_students_excel(self):
        from openpyxl import load_workbook
        upload = self.upload('admission_number,name,father_name,mobile_number\n001,Valid,Father,0123\n002,,Father,0456')
        failed = upload.json()['failed_students']
        response = self.client.post('/super-admin/students/bulk-upload/failed/export',
                                    json={'failed_students': failed})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('failed_students.xlsx', response.headers['content-disposition'])
        book = load_workbook(io.BytesIO(response.content))
        values = list(book.active.values)
        self.assertEqual(values[1][1], '002')
        self.assertEqual(values[1][4], '0456')
        self.assertEqual(values[1][5], 'Invalid student details')
        book.close()
        failed[0]['name'] = '=1+1'
        response = self.client.post('/super-admin/students/bulk-upload/failed/export', json={'failed_students': failed})
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active['C2'].data_type, 's')
        book.close()
        self.user.role = UserRole.ADMIN
        self.assertEqual(self.client.post('/super-admin/students/bulk-upload/failed/export',
            json={'failed_students': failed}).status_code, 403)

    def test_requires_super_admin(self):
        for role in (UserRole.ADMIN, UserRole.SUB_ADMIN, 'Staff', 'Teacher', 'Student'):
            with self.subTest(role=role):
                self.user.role = role
                self.assertEqual(self.upload('admission_number,name,father_name,mobile_number\n001,A,F,123').status_code, 403)
        self.assertEqual(self.db.query(Student).count(), 0)


if __name__ == '__main__':
    unittest.main()
