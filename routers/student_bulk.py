"""Super Admin Excel import; duplicate admissions are skipped."""
import io
import json
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from dependencies.auth import get_current_user
from dependencies.db import get_db_session
from models.school import School
from models.student import Student
from routers.students import validate_student_references, update_student_details
from schemas.student_bulk import StudentBulkRow, bulk_payload, FailedStudentsExportRequest
from utils.dates import today

router = APIRouter()
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000


def failed_student(number, row, reason, errors=None):
    values = row.model_dump() if isinstance(row, StudentBulkRow) else row
    result = {'row': number, 'admission_number': values.get('admission_number') or '',
              'name': values.get('first_name') or values.get('name') or '',
              'father_name': values.get('father_name') or '',
              'mobile_number': values.get('mobile_number') or '', 'reason': reason}
    if errors is not None:
        result['errors'] = errors
    return result


@router.post('/students/bulk-upload', status_code=201)
def bulk_upload_students(
    school_id: int = Query(..., gt=0),
    class_id: int = Query(..., gt=0),
    section_id: int = Query(..., gt=0),
    session_id: int = Query(..., gt=0),
    caste_category_name: str | None = Query(None, min_length=1, max_length=100),
    fee_category_name: str | None = Query(None, min_length=1, max_length=100),
    file: UploadFile = File(...),
    db: Session = Depends(get_db_session),
    user=Depends(get_current_user),
):
    from routers.super_admin import _require_super_admin
    _require_super_admin(user)
    if db.query(School.id).filter_by(id=school_id).first() is None:
        raise HTTPException(404, 'School not found')
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(422, 'Upload an Excel .xlsx file')
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, 'Excel file exceeds 5 MB')
    failed_rows = []
    total = 0
    existing_admissions = {value for (value,) in db.query(Student.admission_number).filter_by(school_id=school_id).all()}
    workbook = None
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise HTTPException(413, 'Expanded Excel file exceeds 50 MB')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        sheet = workbook.worksheets[0]
        # Ignore misleading worksheet dimensions written by third-party exporters.
        sheet.reset_dimensions()
        records = sheet.iter_rows()
        header_cells = next(records, ())
        headers = [str(cell.value).strip() if cell.value is not None else '' for cell in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        required = {'admission_number', 'father_name', 'mobile_number'}
        if not required.issubset(headers) or not {'name', 'first_name'}.intersection(headers):
            raise HTTPException(422, 'Excel requires admission_number, name (or first_name), father_name, mobile_number')
        if not all(headers) or len(headers) != len(set(headers)) or {'name', 'first_name'}.issubset(headers):
            raise HTTPException(422, 'Excel contains blank, duplicate or ambiguous headers')
        rows = []
        for number, cells in enumerate(records, start=2):
            if all(cell.value is None for cell in cells):
                continue
            total += 1
            if total > MAX_ROWS or number > MAX_ROWS + 1:
                raise HTTPException(413, 'Excel exceeds 1000 student rows')
            row = {}
            for index, header in enumerate(headers):
                value = cells[index].value if index < len(cells) else None
                if isinstance(value, int) and not isinstance(value, bool):
                    value = str(value)
                elif isinstance(value, float) and value.is_integer():
                    value = str(int(value))
                row[header] = value
            if any(cell.data_type in ('f', 'e') for cell in cells):
                failed_rows.append(failed_student(number, row, 'Formula and error cells are not supported'))
                continue
            if any(cell.value is not None for cell in cells[len(headers):]):
                failed_rows.append(failed_student(number, row, 'Data found without a column header'))
                continue
            admission = row.get('admission_number')
            admission = admission.strip() if isinstance(admission, str) else admission
            if admission and admission in existing_admissions:
                failed_rows.append(failed_student(number, row, 'Admission number already exists in this school'))
                continue
            try:
                rows.append((number, StudentBulkRow.model_validate(row)))
            except ValidationError as exc:
                failed_rows.append(failed_student(number, row, 'Invalid student details',
                    exc.errors(include_input=False, include_context=False)))
    except HTTPException:
        raise
    except (ParseError, BadZipFile, InvalidFileException, KeyError, ValueError, IndexError, OSError) as exc:
        raise HTTPException(422, 'Upload a valid Excel .xlsx workbook') from exc
    finally:
        if workbook is not None:
            workbook.close()
    if not total:
        raise HTTPException(422, 'Excel must contain at least one student')

    prepared = []
    seen = set()
    for number, row in rows:
        if row.admission_number in seen:
            failed_rows.append(failed_student(number, row, 'Duplicate admission number in Excel'))
            continue
        seen.add(row.admission_number)
        try:
            payload = bulk_payload(row, session_id, class_id, section_id, caste_category_name, fee_category_name)
        except ValidationError as exc:
            failed_rows.append(failed_student(number, row, 'Invalid student details',
                exc.errors(include_input=False, include_context=False)))
            continue
        except ValueError as exc:
            failed_rows.append(failed_student(number, row, str(exc)))
            continue
        try:
            validate_student_references(db, school_id, payload)
        except HTTPException as exc:
            failed_rows.append(failed_student(number, row, exc.detail))
            continue
        prepared.append((number, row, payload))

    try:
        connection = db.connection()
        # SQLite's legacy transaction mode needs BEGIN before savepoints so
        # a final commit failure can roll back every inserted row.
        if connection.dialect.name == 'sqlite' and not connection.connection.driver_connection.in_transaction:
            connection.exec_driver_sql('BEGIN')
        db.query(School).filter_by(id=school_id).with_for_update().one()
        created = []
        for number, row, payload in prepared:
            if db.query(Student.id).filter_by(school_id=school_id, admission_number=row.admission_number).first():
                failed_rows.append(failed_student(number, row, 'Admission number already exists in this school'))
                continue
            try:
                with db.begin_nested():
                    values = payload.student_values()
                    values['admission_date'] = values.get('admission_date') or today()
                    item = Student(school_id=school_id, admission_number=row.admission_number, **values)
                    # Follow the Schools create API for login, addresses and previous school.
                    # The supplied admission number skips automatic admission allocation.
                    update_student_details(db, school_id, item, payload)
                    db.add(item)
                    # Persist shared login IDs before processing the next sibling.
                    db.flush()
            except IntegrityError:
                failed_rows.append(failed_student(number, row, 'Student references conflict with database constraints'))
                continue
            created.append(item)
        db.flush()
        result = [{'id': item.id, 'admission_number': item.admission_number,
                   'name': item.first_name, 'mobile_number': item.mobile_number} for item in created]
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, 'Bulk upload conflicts with existing student records') from exc
    except Exception:
        db.rollback()
        raise
    return {'status': 'success', 'message': 'Students uploaded successfully',
            'school_id': school_id, 'session_id': session_id, 'class_id': class_id,
            'section_id': section_id, 'count': len(result), 'total': total,
            'inserted': len(result), 'failed': len(failed_rows),
            'failed_rows': sorted(failed_rows, key=lambda row: row['row']),
            'failed_students': sorted(failed_rows, key=lambda row: row['row']), 'data': result}


@router.post('/students/bulk-upload/failed/export', response_class=Response,
             responses={200: {'content': {'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': {}}}})
def export_failed_students(payload: FailedStudentsExportRequest, user=Depends(get_current_user)):
    """Download the failed_students list returned by the bulk upload as Excel."""
    from routers.super_admin import _require_super_admin
    _require_super_admin(user)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Failed Students'
    columns = ['row', 'admission_number', 'name', 'father_name', 'mobile_number', 'reason', 'errors']
    sheet.append(columns)
    for student in payload.failed_students:
        values = student.model_dump()
        values['errors'] = json.dumps(values['errors'], ensure_ascii=False) if values['errors'] else ''
        sheet.append([values[key] for key in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Export user input as literal text, including strings starting with '='.
                cell.data_type = 's'
                cell.number_format = '@'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='215E80')
    for column, width in zip('ABCDEFG', (12, 24, 30, 30, 24, 60, 90)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return Response(stream.getvalue(),
                    media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename="failed_students.xlsx"',
                             'Cache-Control': 'no-store'})
